"""Audit supplied workbook values without asserting historical model provenance."""
from collections import Counter
from pathlib import Path
import hashlib
import json
import math
import sys

import openpyxl


def audit(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = {'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheets': {}}
    for name in ['Sinle_OFE', 'Two_OFE']:
        sheet = book[name]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        indices = {v: i for i, v in enumerate(header) if v}
        records = []
        for row in rows:
            if isinstance(row[indices['OFE']], (int, float)) and isinstance(row[indices['Y']], (int, float)):
                record = {k: row[i] for k, i in indices.items()}
                if not all(isinstance(v, (int, float)) and math.isfinite(v) for v in record.values()):
                    raise ValueError(f'Nonfinite/non-numeric data in {name}')
                records.append(record)
        counts = Counter((r['Y'], r['J'], r['OFE']) for r in records)
        if any(v != 1 for v in counts.values()):
            raise ValueError(f'Duplicate daily OFE rows in {name}')
        ofes = sorted(set(r['OFE'] for r in records))
        areas = {int(o): sorted(set(r['Area'] for r in records if r['OFE'] == o)) for o in ofes}
        if any(len(a) != 1 for a in areas.values()):
            raise ValueError('Changing area requires separate interval treatment')
        area = sum(a[0] for a in areas.values())
        years = sorted(set(r['Y'] for r in records))
        sums = {k: sum(r[k] * r['Area'] / area for r in records) / len(years)
                for k in ['P', 'Ep', 'Es', 'Er', 'Dp']}
        terminal = [r for r in records if r['OFE'] == max(ofes)]
        sums['terminal_Q_mm_yr'] = sum(r['Q'] for r in terminal) / len(years)
        sums['terminal_lateral_hillslope_mm_yr'] = sum(r['latqcc'] * r['Area'] / area for r in terminal) / len(years)
        sums['P_minus_ET_Dp_Q_lateral_mm_yr'] = sums['P'] - sum(sums[k] for k in ['Ep', 'Es', 'Er', 'Dp', 'terminal_Q_mm_yr', 'terminal_lateral_hillslope_mm_yr'])
        result['sheets'][name] = {'rows': len(records), 'years': years, 'areas_m2': areas,
            'terminal_fraction': areas[int(max(ofes))][0] / area, 'annual': sums,
            'disposition': 'accounting screen only; initial storage, historical executable and inputs unavailable'}
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    result['summary_formulas'] = {c.coordinate: c.value for row in formulas.worksheets[0] for c in row if c.data_type == 'f'}
    return result


if __name__ == '__main__':
    Path(sys.argv[2]).write_text(json.dumps(audit(Path(sys.argv[1])), indent=2) + '\n')
